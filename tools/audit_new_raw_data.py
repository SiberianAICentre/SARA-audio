"""Read-only corpus audit; write an inventory and joins without changing inputs."""

from __future__ import annotations

import ast
import csv
import hashlib
import json
import re
import wave
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data/new_raw_data"
OUT = ROOT / "NEW_RAW_DATA_INVENTORY.json"
ID = re.compile(r"(\d{14})_(?:Agent_|Оператор)(\d+)")
LINE = re.compile(r"^\[(\d+):(\d{2})\]\s+([^:]+):\s?(.*)$")


def key(value: str) -> str | None:
    match = ID.search(value)
    return f"{match[1]}_Agent_{int(match[2]):02d}" if match else None


def sha(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def table(path: Path) -> list[dict]:
    delimiter = "\t" if path.suffix == ".tsv" else ";" if "custom" in path.stem else ","
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream, delimiter=delimiter))


def wav_info(stream) -> dict:
    try:
        with wave.open(stream, "rb") as audio:
            return dict(
                duration_s=audio.getnframes() / audio.getframerate(),
                sample_rate=audio.getframerate(),
                channels=audio.getnchannels(),
                sample_width=audio.getsampwidth(),
            )
    except (wave.Error, EOFError) as error:
        return {"error": str(error)}


def main() -> None:
    inventory, archives, tables, workbooks = [], [], [], []
    records = defaultdict(lambda: defaultdict(list))
    texts = {}
    for path in sorted(DATA.rglob("*")):
        if not path.is_file():
            continue
        relative = path.relative_to(DATA).as_posix()
        item = dict(path=relative, bytes=path.stat().st_size, sha256=sha(path))
        if path.name.startswith("._") or path.name == ".DS_Store":
            item["kind"] = "macos_metadata"
            inventory.append(item)
            continue
        record_id = key(path.name)
        if path.suffix == ".wav":
            with path.open("rb") as stream:
                item.update(wav_info(stream))
            records[record_id]["audio_disk"].append(item)
        elif path.suffix == ".txt" and record_id:
            text = path.read_text(encoding="utf-8-sig")
            lines = text.splitlines()
            parsed = [LINE.match(line) for line in lines]
            stamps = [int(m[1]) * 60 + int(m[2]) for m in parsed if m]
            item.update(
                lines=len(lines),
                unmatched_lines=[i + 1 for i, m in enumerate(parsed) if not m],
                last_timestamp_s=max(stamps, default=None),
                descending_timestamps=sum(b < a for a, b in zip(stamps, stamps[1:], strict=False)),
                repeated_timestamps=sum(b == a for a, b in zip(stamps, stamps[1:], strict=False)),
            )
            variant = (
                "processed" if relative.startswith("Обработанные") else "before_roles"
            )
            records[record_id][variant].append(item)
            texts[(variant, record_id)] = (path, parsed)
        elif path.suffix == ".zip":
            with zipfile.ZipFile(path) as archive:
                members = []
                for entry in archive.infolist():
                    if entry.is_dir():
                        continue
                    member = dict(
                        path=entry.filename,
                        bytes=entry.file_size,
                        crc32=f"{entry.CRC:08x}",
                    )
                    if (
                        entry.filename.lower().endswith(".wav")
                        and not Path(entry.filename).name.startswith("._")
                        and "__MACOSX" not in entry.filename
                    ):
                        with archive.open(entry) as stream:
                            member.update(wav_info(stream))
                        records[key(entry.filename)]["audio_zip"].append(
                            dict(archive=relative, **member)
                        )
                    members.append(member)
                archives.append(
                    dict(
                        path=relative,
                        files=len(members),
                        extensions=dict(
                            Counter(Path(m["path"]).suffix for m in members)
                        ),
                        members=members,
                    )
                )
        elif path.suffix in {".csv", ".tsv"}:
            rows = table(path)
            tables.append(
                dict(
                    path=relative, rows=len(rows), columns=list(rows[0]) if rows else []
                )
            )
            if path.suffix == ".csv":
                for index, row in enumerate(rows, 2):
                    record_id = key(row.get("original_file_path", ""))
                    records[record_id]["tables"].append(
                        dict(
                            path=relative,
                            row=index,
                            duration_sec=row.get("duration_sec"),
                            word_count=row.get("word_count"),
                            text_chars=len(row.get("text", "")),
                            agent_id=row.get("agent_id"),
                        )
                    )
        elif path.suffix == ".xlsx":
            if not zipfile.is_zipfile(path):
                workbooks.append(
                    dict(
                        path=relative,
                        error="Not an OOXML ZIP workbook",
                        magic_hex=path.open("rb").read(16).hex(),
                    )
                )
                inventory.append(item)
                continue
            book = load_workbook(path, read_only=True, data_only=True)
            sheets = []
            for sheet in book:
                rows = list(sheet.iter_rows(values_only=True))
                sheets.append(
                    dict(
                        name=sheet.title,
                        rows=sheet.max_row,
                        columns=sheet.max_column,
                        first_rows=[list(row) for row in rows[:3]],
                    )
                )
            book.close()
            workbooks.append(dict(path=relative, sheets=sheets))
        inventory.append(item)

    processed = DATA / "Обработанные_записи (1)"
    registry = table(processed / "_Проверка/Реестр.tsv")
    edits = table(processed / "_Проверка/Замены_тишины.tsv")
    roles = table(processed / "_Проверка/Назначение_ролей.tsv")
    edits_by_line = {(row["Исходный путь"], int(row["Строка"])): row for row in edits}
    checks = Counter()
    issues = []
    for row in registry:
        relative = row["Исходный путь"]
        record_id = key(relative)
        before_path, before = texts[("before_roles", record_id)]
        after_path, after = texts[("processed", record_id)]
        checks["source_hash_matches"] += sha(before_path) == row["SHA256 исходника"]
        checks["result_hash_matches"] += sha(after_path) == row["SHA256 результата"]
        checks["line_counts_match"] += (
            len(before) == len(after) == int(row["Строк с репликами"])
        )
        records[record_id]["review"].append(
            {
                k: row[k]
                for k in ["Точечных назначений роли", "Замен артефактов", "Примечаний"]
            }
        )
        for n, (a, b) in enumerate(zip(before, after, strict=True), 1):
            if not a or not b:
                issues.append(dict(id=record_id, line=n, error="unparsed"))
                continue
            checks["timestamps_preserved"] += (a[1], a[2]) == (b[1], b[2])
            checks["valid_roles"] += b[3] in {"Оператор", "Клиент"}
            edit = edits_by_line.get((relative, n))
            if a[4] == b[4]:
                checks["unchanged_text_lines"] += 1
            elif (
                edit
                and a[4] == edit["Исходный текст"]
                and b[4] == edit["Итоговый текст"]
            ):
                checks["documented_text_changes"] += 1
            else:
                issues.append(
                    dict(id=record_id, line=n, error="undocumented_text_change")
                )
        disk = records[record_id].get("audio_disk", [])
        zip_audio = records[record_id].get("audio_zip", [])
        audio = disk or zip_audio
        if audio and "duration_s" in audio[0]:
            records[record_id]["alignment_check"] = dict(
                audio_duration_s=audio[0]["duration_s"],
                last_text_timestamp_s=max(
                    int(m[1]) * 60 + int(m[2]) for m in after if m
                ),
            )

    unique_sets = {
        kind: {k for k, r in records.items() if r.get(kind)}
        for kind in ("audio_disk", "audio_zip", "before_roles", "processed")
    }
    audio_ids = unique_sets["audio_disk"] | unique_sets["audio_zip"]
    summary = dict(
        files=len(inventory),
        extensions=dict(Counter(Path(i["path"]).suffix for i in inventory)),
        unique_ids={k: len(v) for k, v in unique_sets.items()},
        audio_unique_total=len(audio_ids),
        text_without_audio=sorted(unique_sets["processed"] - audio_ids),
        audio_without_text=sorted(audio_ids - unique_sets["processed"]),
        registry_checks=dict(checks),
        registry_issues=issues,
        registry_rows=len(registry),
        role_rows=len(roles),
        replacement_rows=len(edits),
        timestamps_beyond_audio=sum(
            r.get("alignment_check", {}).get("last_text_timestamp_s", 0)
            > r.get("alignment_check", {}).get("audio_duration_s", float("inf"))
            for r in records.values()
        ),
    )
    payload = dict(
        root="data/new_raw_data",
        summary=summary,
        files=inventory,
        archives=archives,
        tables=tables,
        workbooks=workbooks,
        records=dict(sorted(records.items(), key=lambda x: str(x[0]))),
    )
    full = table(DATA / "2_file/df_full_standard.csv")
    train = table(DATA / "2_file/df_train_standard.csv")
    test = table(DATA / "2_file/df_test_standard.csv")
    train_ids = {key(r["original_file_path"]) for r in train}
    test_ids = {key(r["original_file_path"]) for r in test}
    relations = dict(
        train_test_overlap=sorted(train_ids & test_ids),
        full_equals_train_plus_test=full == train + test,
        full_ids_order_equals_train_plus_test=[r["original_file_path"] for r in full]
        == [r["original_file_path"] for r in train + test],
        train_test_agent_overlap=sorted(
            {r["agent_id"] for r in train} & {r["agent_id"] for r in test}
        ),
        table_audio_duration_matches=0,
        table_audio_duration_max_error_s=0,
        audio_total_duration_s=0,
        formats=dict(
            Counter(
                str((a.get("sample_rate"), a.get("channels"), a.get("sample_width")))
                for r in records.values()
                for a in (r.get("audio_disk") or r.get("audio_zip") or [])[:1]
            )
        ),
        csv_variants_equal={
            split: table(DATA / f"2_file/df_{split}_custom.csv")
            == table(DATA / f"2_file/df_{split}_standard.csv")
            for split in ("full", "train", "test")
        },
    )
    full_by_id = {row["original_file_path"]: row for row in full}
    subset_by_id = {row["original_file_path"]: row for row in train + test}
    relations["full_vs_subsets_different_fields"] = dict(
        Counter(
            field
            for name, row in full_by_id.items()
            for field in row
            if row[field] != subset_by_id[name][field]
        )
    )
    relations["pause_fields_numeric_equal"] = all(
        float(row[field]) == float(subset_by_id[name][field])
        for name, row in full_by_id.items()
        for field in ("avg_interword_pause_ms", "std_interword_pause_ms")
    )
    book = load_workbook(
        DATA / "Анализ_транскрипций_НМ.xlsx", read_only=True, data_only=True
    )
    expert_rows = list(book.worksheets[0].values)[1:]
    relations["expert_text_equals_full_same_order"] = len(expert_rows) == len(
        full
    ) and all(a[1] == b["text"] for a, b in zip(expert_rows, full, strict=True))
    book.close()
    for index, row in enumerate(full, 2):
        record_id = key(row["original_file_path"])
        record = records[record_id]
        record["split"] = "train" if record_id in train_ids else "test"
        record["expert_sheet_row"] = (
            index if relations["expert_text_equals_full_same_order"] else None
        )
        audio = (record.get("audio_disk") or record.get("audio_zip"))[0]
        error = abs(float(row["duration_sec"]) - audio["duration_s"])
        relations["table_audio_duration_matches"] += error < 0.00001
        relations["table_audio_duration_max_error_s"] = max(
            relations["table_audio_duration_max_error_s"], error
        )
        relations["audio_total_duration_s"] += audio["duration_s"]
    payload["relations"] = relations
    feature_registry = json.loads(
        (ROOT / "src/sara_audio/data/linguistic_features.json").read_text(
            encoding="utf-8"
        )
    )
    tree = ast.parse(
        (ROOT / "src/sara_audio/linguistic.py").read_text(encoding="utf-8")
    )
    rules = next(
        node.value
        for node in ast.walk(tree)
        if isinstance(node, ast.AnnAssign)
        and isinstance(node.target, ast.Attribute)
        and node.target.attr == "_rules"
    )
    implemented = {k.value for k in rules.keys}
    payload["linguistic_registry"] = [
        {**feature, "rule_implemented": feature["code"] in implemented}
        for feature in feature_registry["features"]
    ]
    summary["macos_metadata_files"] = sum(
        i.get("kind") == "macos_metadata" for i in inventory
    )
    summary["regular_data_files"] = len(inventory) - summary["macos_metadata_files"]
    OUT.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    md = [
        "# Пофайловая карта новых данных",
        "",
        "Сформировано скриптом tools/audit_new_raw_data.py. Пути исходников относительны data/new_raw_data. ZIP указан через `archive::member`; архивы не распаковывались.",
        "",
        "| ID | Split | WAV (диск или ZIP) | До назначения ролей | Обработанный текст | Строка Excel НМ | WAV, с | Последняя отметка текста, с |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for record_id, record in sorted(records.items(), key=lambda x: str(x[0])):
        audio = record.get("audio_disk") or record.get("audio_zip") or []
        location = (
            (
                (audio[0].get("archive", "") + "::" if "archive" in audio[0] else "")
                + audio[0]["path"]
            )
            if audio
            else "не найдено"
        )
        before = "; ".join(t["path"] for t in record.get("before_roles", []))
        after = "; ".join(t["path"] for t in record.get("processed", []))
        md.append(
            f"| {record_id} | {record.get('split', 'text_only')} | {location} | {before} | {after} | {record.get('expert_sheet_row', '')} | {audio[0]['duration_s'] if audio else ''} | {record.get('processed', [{}])[0].get('last_timestamp_s', '')} |"
        )
    (ROOT / "NEW_RAW_DATA_RECORD_MAP.md").write_text(
        "\n".join(md) + "\n", encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=True, indent=2))
    print(json.dumps(relations, ensure_ascii=True, indent=2))
    print("Inventory:", OUT.name)


if __name__ == "__main__":
    main()
